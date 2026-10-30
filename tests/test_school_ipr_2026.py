"""Anonymized examples of the current teacher DOCX layout."""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement

from ircenter_automation.school_ipr_2026 import _ENVIRONMENT_ROWS, extract_portal_fields


@pytest.mark.parametrize("section9_choices", [
    (" Так", " Ні"),
    ("+ Так", " Ні"),
    (" Так", "+ Ні"),
])
def test_extracts_matching_sections_and_skips_section9(tmp_path: Path, section9_choices):
    doc = Document()
    for text in (
        "3. Категорії (типи) особливих освітніх потреб: мовленнєві труднощі",
        "4. Рівень підтримки (протокол засідання від 12 жовтня 2026 р. № 7):",
        " перший рівень підтримки", "+ другий рівень підтримки",
        " третій рівень підтримки", " четвертий рівень підтримки",
        " п’ятий рівень підтримки",
        "5. Створення інклюзивного освітнього середовища",
        "наявність індивідуального навчального плану:", " так", "+ ні",
        "здійснення освітнього процесу:", "+ за розкладом класу", " за адаптованим розкладом",
        "забезпечення допоміжними засобами для навчання:", " так", "+ ні",
        "8. Особливості процесу оцінювання:",
        "+ спрощення інструкцій", "+ гнучкість використання часу",
        "+ зменшення кількості завдань", "+ використання оптимальних форм оцінювання",
        "+ використання доступних форм викладу інформації", "+ використання допоміжних засобів",
        " уникнення графічних завдань", " уникнення зорових/слухових подразників",
        " інше __________________",
        "9. Потреба у продовженні строків навчання:", *section9_choices,
        "прізвище, власне ім’я, по батькові (за наявності) батьків (інших законних представників) учня Шевченко Олена Іванівна",
        "Я, Шевченко Олена Іванівна,",
    ):
        doc.add_paragraph(text)
    subjects = doc.add_table(rows=1, cols=3)
    subjects.rows[0].cells[0].text = "Найменування навчального предмета"
    for name, adaptation in (
        ("Математика", "Адаптація форми усної відповіді за потреби; додатковий час. "
         "Модифікація змісту й очікуваних результатів не застосовується."),
        ("Читання", "+ Модифікація"),
    ):
        cells = subjects.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = name, adaptation, "Учитель Приклад"
    goals = doc.add_table(rows=1, cols=5)
    goals.rows[0].cells[0].text = "Освітня галузь/ навчальний предмет"
    cells = goals.add_row().cells
    for cell, value in zip(cells, ("Мовна", "Читає", "Читати", "Прочитає", "")):
        cell.text = value
    path = tmp_path / "example.docx"
    doc.save(path)
    base = {"execution_period": "з 01.09.2026 по 31.05.2027", "year_of_study": "2",
            "child_grade": "2-А", "parents_phones": "000 000 000"}
    result = {field.name: field.value for field in extract_portal_fields(path, base)}
    assert result["3[1][1]"] == "мовленнєві труднощі"
    assert result["supportLevel"] == "2"
    assert result["5[26][1]"] == "0"
    assert result["5[27][1]"] == "за розкладом класу"
    assert result["5[29][1]"] == "Математика"
    assert result["5[29][2]"] == "Адаптація"
    assert result["5[30][1]"] == "Читання"
    assert result["5[30][2]"] == "Модифікація"
    assert result["6[1][4]"] == "Прочитає"
    assert result["8[7][1]"] == "false"
    assert result["8[9][0]"] == "false"
    assert "9[1][1]" not in result
    assert result["footer_parents_names"] == "Шевченко Олена Іванівна"


def test_subject_choice_does_not_guess_when_both_methods_are_used():
    from ircenter_automation.school_ipr_2026 import _subject_choice

    value = "Адаптація усної відповіді. Модифікація змісту застосовується."
    assert _subject_choice(value) == value


@pytest.mark.parametrize("choices, expected", [
    (("так", "ні"), "0"),
    ((" так", " ні"), "0"),
    (("+ так", " ні"), "1"),
    ((" так", "+ ні"), "0"),
    (("+ так", "ні"), "1"),
])
def test_individual_plan_defaults_to_no_only_without_explicit_selection(tmp_path: Path, choices, expected):
    doc = Document()
    doc.add_paragraph("5. Створення інклюзивного освітнього середовища")
    doc.add_paragraph("наявність індивідуального навчального плану:")
    for choice in choices:
        doc.add_paragraph(choice)
    path = tmp_path / "plan.docx"
    doc.save(path)
    base = {"execution_period": "з 01.09.2026 по 31.05.2027", "year_of_study": "2",
            "child_grade": "2-А", "parents_phones": "000 000 000"}
    result = {field.name: field.value for field in extract_portal_fields(path, base)}
    assert result["5[26][1]"] == expected


@pytest.mark.parametrize("underlined_option, expected", [
    (0, "за розкладом класу"),
    (1, "за адаптованим розкладом з відвідуванням окремих уроків за індивідуальним розкладом відповідно до індивідуального навчального плану"),
    (None, None),
    ("both", None),
])
def test_underlined_schedule_option_counts_as_selected(tmp_path: Path, underlined_option, expected):
    doc = Document()
    doc.add_paragraph("5. Створення інклюзивного освітнього середовища")
    doc.add_paragraph("здійснення освітнього процесу:")
    options = (
        "за розкладом класу",
        "за адаптованим розкладом з відвідуванням окремих уроків за індивідуальним розкладом відповідно до індивідуального навчального плану",
    )
    for index, option in enumerate(options):
        paragraph = doc.add_paragraph()
        box = paragraph.add_run("  ")
        box.underline = underlined_option is None and index == 0
        label = paragraph.add_run(option)
        label.underline = index == underlined_option or underlined_option == "both"
    path = tmp_path / "schedule.docx"
    doc.save(path)
    base = {"execution_period": "з 01.09.2026 по 31.05.2027", "year_of_study": "2",
            "child_grade": "2-А", "parents_phones": "000 000 000"}
    result = {field.name: field.value for field in extract_portal_fields(path, base)}
    assert result.get("5[27][1]") == expected


def test_underlined_plan_yes_overrides_default_no(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("5. Створення інклюзивного освітнього середовища")
    doc.add_paragraph("наявність індивідуального навчального плану:")
    paragraph = doc.add_paragraph(" ")
    paragraph.add_run("так").underline = True
    doc.add_paragraph(" ні")
    path = tmp_path / "plan.docx"
    doc.save(path)
    base = {"execution_period": "з 01.09.2026 по 31.05.2027", "year_of_study": "2",
            "child_grade": "2-А", "parents_phones": "000 000 000"}
    result = {field.name: field.value for field in extract_portal_fields(path, base)}
    assert result["5[26][1]"] == "1"


def test_underlined_choices_in_capitalized_school_template(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Адреса задекларованого/зареєстрованого місця проживання (перебування): вул. Прикладна, 1")
    doc.add_paragraph("Дата зарахування до закладу загальної середньої освіти: 01.09.2026")
    doc.add_paragraph("Прізвище, власне ім’я, по батькові (за наявності) батьків: Прикладна Олена Іванівна")
    doc.add_paragraph("Я, Прикладна Олена Іванівна")
    doc.add_paragraph("4. Рівень підтримки:")
    for index, ordinal in enumerate(("перший", "другий", "третій", "четвертий", "п’ятий")):
        doc.add_paragraph().add_run(f"{ordinal} рівень підтримки").underline = index == 3
    doc.add_paragraph("5. Створення інклюзивного освітнього середовища")
    doc.add_paragraph("Наявність індивідуального навчального плану:")
    doc.add_paragraph("так")
    doc.add_paragraph().add_run("ні").underline = True
    doc.add_paragraph("Здійснення освітнього процесу:")
    doc.add_paragraph().add_run("за розкладом класу").underline = True
    doc.add_paragraph("за адаптованим розкладом з відвідуванням окремих уроків за індивідуальним розкладом відповідно до індивідуального навчального плану")
    doc.add_paragraph("Забезпечення допоміжними засобами для навчання:")
    doc.add_paragraph("так")
    doc.add_paragraph().add_run("ні").underline = True
    doc.add_paragraph("8. Особливості процесу оцінювання:")
    labels = (
        "спрощення інструкцій", "гнучкість використання часу",
        "зменшення кількості завдань", "використання оптимальних форм оцінювання",
        "використання доступних форм викладу інформації",
        "використання допоміжних засобів", "уникнення графічних завдань",
        "уникнення зорових/слухових подразників",
    )
    for index, label in enumerate(labels):
        doc.add_paragraph().add_run(label).underline = index < 4
    doc.add_paragraph("інше __________________")
    doc.add_paragraph("9. Потреба у продовженні строків навчання:")
    path = tmp_path / "underlined.docx"
    doc.save(path)
    base = {"execution_period": "з 01.09.2026 по 31.05.2027", "year_of_study": "4",
            "child_grade": "6-А", "parents_phones": "+380(00-000-00-00)"}
    result = {field.name: field.value for field in extract_portal_fields(path, base)}
    assert result["parent_address"] == "вул. Прикладна, 1"
    assert result["child_enrollment_date"] == "01.09.2026"
    assert result["footer_parents_names"] == "Прикладна Олена Іванівна"
    assert result["supportLevel"] == "4"
    assert result["5[26][1]"] == "0"
    assert result["5[27][1]"] == "за розкладом класу"
    assert result["5[28][1]"] == "0"
    assert [result[f"8[{i}][1]"] for i in range(1, 9)] == ["true"] * 4 + ["false"] * 4
    assert result["8[9][0]"] == "false"


def test_reads_tables_inside_word_content_controls_and_skips_heading_only_goal_row(tmp_path: Path):
    doc = Document()

    def wrap(table):
        element = table._element
        parent = element.getparent()
        control = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        control.append(content)
        parent.replace(element, control)
        content.append(element)

    environment = doc.add_table(rows=1, cols=2)
    environment.cell(0, 0).text = "Назва адаптації/модифікації освітнього середовища"
    for label, _number in _ENVIRONMENT_ROWS:
        row = environment.add_row()
        row.cells[0].text, row.cells[1].text = label, "Потрібний захід"
    wrap(environment)

    subjects = doc.add_table(rows=1, cols=3)
    subjects.cell(0, 0).text = "Найменування навчального предмета"
    row = subjects.add_row()
    row.cells[0].text, row.cells[1].text, row.cells[2].text = "Математика", "Адаптація", "Учитель Приклад"
    wrap(subjects)

    goals = doc.add_table(rows=1, cols=5)
    goals.cell(0, 0).text = "Освітня галузь/ навчальний предмет"
    for values in (("Мовна", "Читає", "Навчитися", "Прочитає"),
                   ("Підзаголовок", "", "", ""),
                   ("Математична", "Рахує", "Навчитися", "Порахує")):
        row = goals.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = value
    wrap(goals)

    services = doc.add_table(rows=1, cols=4)
    services.cell(0, 0).text = "Найменування заняття"
    row = services.add_row()
    for cell, value in zip(row.cells, ("Заняття", "Фахівець Приклад", "Вівторок", "10:00")):
        cell.text = value
    wrap(services)

    team = doc.add_table(rows=1, cols=2)
    team.cell(0, 0).text = "Прізвище, власне ім’я, по батькові"
    row = team.add_row()
    row.cells[0].text, row.cells[1].text = "Фахівець Приклад", "Учитель"
    wrap(team)

    path = tmp_path / "wrapped-tables.docx"
    doc.save(path)
    assert not Document(path).tables
    base = {"execution_period": "з 01.09.2026 по 31.05.2027", "year_of_study": "4",
            "child_grade": "6-А", "parents_phones": "000 000 000"}
    result = {field.name: field.value for field in extract_portal_fields(path, base)}
    assert result["5[1][1]"] == "Потрібний захід"
    assert result["5[29][1]"] == "Математика"
    assert result["5[29][2]"] == "Адаптація"
    assert result["6[1][1]"] == "Мовна"
    assert result["6[2][1]"] == "Математична"
    assert "6[3][1]" not in result
    assert result["7[1][1]"] == "Заняття"
    assert result["10[1][2]"] == "Учитель"


@pytest.mark.parametrize("markers, expected", [
    (("+ ",) * 5 + ("• ",) * 3, ["true"] * 5 + ["false"] * 3),
    (("+ ",) * 5 + ("",) * 3, ["true"] * 5 + ["false"] * 3),
    (("+ • ", "• + ", "+ • ", "• ", "• ", "•+ ", "• ", "• "),
     ["true"] * 3 + ["false"] * 2 + ["true"] + ["false"] * 2),
])
def test_assessment_bullets_and_mixed_plus_markers(tmp_path: Path, markers, expected):
    doc = Document()
    doc.add_paragraph("8. Особливості процесу оцінювання:")
    labels = (
        "спрощення інструкцій", "гнучкість використання часу",
        "зменшення кількості завдань", "використання оптимальних форм оцінювання",
        "використання доступних форм викладу інформації",
        "використання допоміжних засобів", "уникнення графічних завдань",
        "уникнення зорових/слухових подразників",
    )
    for marker, label in zip(markers, labels):
        doc.add_paragraph(marker + label)
    doc.add_paragraph("• інше __________________")
    doc.add_paragraph("9. Потреба у продовженні строків навчання:")
    path = tmp_path / "assessment.docx"
    doc.save(path)
    base = {"execution_period": "з 01.09.2026 по 31.05.2027", "year_of_study": "2",
            "child_grade": "2-А", "parents_phones": "000 000 000"}
    result = {field.name: field.value for field in extract_portal_fields(path, base)}
    assert [result[f"8[{index}][1]"] for index in range(1, 9)] == expected
    assert result["8[9][0]"] == "false"


def test_section_10_continues_across_multiple_word_tables(tmp_path: Path):
    doc = Document()
    for member, role in (("Перший Учитель", "Вчитель"), ("Другий Учитель", "Асистент")):
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Прізвище, власне ім’я, по батькові (за наявності)"
        table.cell(0, 1).text = "Найменування посади"
        row = table.add_row()
        row.cells[0].text, row.cells[1].text = member, role
    path = tmp_path / "split-team.docx"
    doc.save(path)
    base = {"execution_period": "01.09.2026-28.05.2027", "year_of_study": "2",
            "child_grade": "7-Б", "parents_phones": "0000000000"}
    fields = {field.name: field.value for field in extract_portal_fields(path, base)}
    assert fields["10[1][1]"] == "Перший Учитель"
    assert fields["10[1][2]"] == "Вчитель"
    assert fields["10[2][1]"] == "Другий Учитель"
    assert fields["10[2][2]"] == "Асистент"


def test_section_10_skips_repeated_header_inside_continuation(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Прізвище, власне ім’я, по батькові"
    table.cell(0, 1).text = "Найменування посади"
    repeated = table.add_row()
    repeated.cells[0].text = "Прізвище, власне ім’я, по батькові"
    repeated.cells[1].text = "Найменування посади"
    member = table.add_row()
    member.cells[0].text = "Третій Учитель"
    member.cells[1].text = "Вчитель"
    path = tmp_path / "repeated-team-header.docx"
    doc.save(path)
    base = {"execution_period": "01.09.2026-28.05.2027", "year_of_study": "2",
            "child_grade": "7-Б", "parents_phones": "0000000000"}
    fields = {field.name: field.value for field in extract_portal_fields(path, base)}
    assert fields["10[1][1]"] == "Третій Учитель"
    assert "10[2][1]" not in fields
