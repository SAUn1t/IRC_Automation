from pathlib import Path

import pytest
from docx import Document

from ircenter_automation.browser import require_selectors
from ircenter_automation.documents import find_document, inspect_document, parse_document
from ircenter_automation.errors import (
    ChildIdentityMismatch, ConfigurationError, DocumentNotFound, DocumentParsingError,
    MultipleDocumentsFound, WorkflowError,
)
from ircenter_automation.mapping import FieldSpec, map_document
from ircenter_automation.normalization import (
    comparison_key, labeled_school_full_name, parse_child_name, resolve_two_word_name,
)
from ircenter_automation.state import State, StateMachine


def test_unicode_name_normalization_and_exactness():
    child = parse_child_name("  Іва́н   О’Коннор-Шевченко  ")
    assert child.display == "Іва́н О’Коннор-Шевченко"
    assert child.matches("ІВА́Н", "О’КОННОР-ШЕВЧЕНКО")
    assert not child.matches("Іван", "О’Коннор-Шевченко")
    assert not child.matches("Іва́н", "О'Коннор-Шевченко")
    with pytest.raises(WorkflowError):
        parse_child_name("Іван Іванович Шевченко")


def test_document_discovery_stops_on_missing_or_multiple(tmp_path: Path):
    child = parse_child_name("Іван Шевченко")
    patterns = ("{Surname} {Name}.docx", "{Name} {Surname}.docx")
    with pytest.raises(DocumentNotFound):
        find_document(tmp_path, child, patterns)
    first = tmp_path / "Шевченко Іван.docx"
    first.touch()
    assert find_document(tmp_path, child, patterns) == first
    (tmp_path / "Іван Шевченко.docx").touch()
    with pytest.raises(MultipleDocumentsFound):
        find_document(tmp_path, child, patterns)


def test_school_name_resolves_either_two_word_order_from_labeled_full_name(tmp_path: Path):
    full_name = "Шевченко Іван Петрович"
    for label in (
        "прізвище, власне ім’я, по батькові(за наявності) ",
        "прізвище, власне ім’я, по батькові (за наявності) ",
        "прізвище, власне ім’я, по батькові\u00a0",
        "Прізвище, власне ім’я, по батькові (за наявності): ",
    ):
        assert labeled_school_full_name(label + full_name) == full_name
    assert labeled_school_full_name(
        "прізвище, власне ім’я, по батькові батьків Шевченко Олена Іванівна"
    ) is None
    for entered in ("Шевченко Іван", "Іван Шевченко"):
        assert resolve_two_word_name(parse_child_name(entered), full_name).display == "Іван Шевченко"
    with pytest.raises(ChildIdentityMismatch):
        resolve_two_word_name(parse_child_name("Іван Петренко"), full_name)
    path = tmp_path / "Шевченко Іван.docx"
    path.touch()
    assert find_document(tmp_path, parse_child_name("Шевченко Іван"), ("{Surname} {Name}.docx",)) == path
    assert find_document(tmp_path, parse_child_name("Іван Шевченко"), ("{Surname} {Name}.docx",)) == path


def test_docx_paragraph_table_empty_and_multiline(tmp_path: Path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Анкета", 1)
    doc.add_paragraph("Дата народження: 01.02.2015")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Примітки"
    table.cell(0, 1).text = "Рядок 1\nРядок 2"
    table.cell(1, 0).text = "Згода"
    table.cell(1, 1).text = "☐"
    doc.save(path)
    parsed = parse_document(path)
    assert parsed.headings == ("Анкета",)
    assert parsed.fields[comparison_key("Примітки")] == "Рядок 1\nРядок 2"
    assert parsed.fields[comparison_key("Згода")] == "no"


def test_docx_duplicate_and_malformed_stop(tmp_path: Path):
    path = tmp_path / "duplicate.docx"
    doc = Document()
    doc.add_paragraph("Вік: 9")
    doc.add_paragraph("Вік: 10")
    doc.save(path)
    with pytest.raises(DocumentParsingError):
        parse_document(path)
    corrupt = tmp_path / "corrupt.docx"
    corrupt.write_bytes(b"not a Word document")
    with pytest.raises(DocumentParsingError):
        parse_document(corrupt)


@pytest.mark.parametrize("class_line, expected_grade", (
    ("клас навчання 2-А", "2-А"),
    ("Клас навчання 2-А", "2-А"),
    ("Клас навчання: 2-А", "2-А"),
    ("Клас навчання:2-А", "2-А"),
    ("Клас навчання: 6-А", "6-А"),
))
@pytest.mark.parametrize("year_line, expected_year", (
    ("рік навчання 2", "2"),
    ("Рік навчання 2", "2"),
    ("Рік навчання: 2", "2"),
    ("Рік навчання:2", "2"),
    ("Рік навчання: 4-й", "4"),
))
def test_school_ipr_extracts_only_exact_scalar_counterparts(
    tmp_path: Path, class_line: str, expected_grade: str, year_line: str, expected_year: str,
):
    path = tmp_path / "synthetic.docx"
    doc = Document()
    for line in (
        "1. Період виконання з 01.09.2026 по 31.05.2027",
        "прізвище, власне ім’я, по батькові\u00a0Шевченко Іван Петрович",
        "контактні номери телефонів батьків (інших законних представників) учня 000 000 000",
        class_line, year_line,
        "5. Створення інклюзивного освітнього середовища, зокрема:",
        "11. Погодження індивідуальної програми розвитку з батьками",
    ):
        doc.add_paragraph(line)
    doc.save(path)
    assert inspect_document(path).startswith("11-section school IPR")
    parsed = parse_document(path)
    assert parsed.template == "school_ipr_2026"
    assert parsed.fields["execution_period"] == "з 01.09.2026 по 31.05.2027"
    assert parsed.fields["child_grade"] == expected_grade
    assert parsed.fields["year_of_study"] == expected_year
    assert len(parsed.fields) == 5  # Four transferable fields plus the child identity guard.


def test_school_ipr_reads_underlined_phone_after_uchenytsi_colon(tmp_path: Path):
    path = tmp_path / "synthetic.docx"
    doc = Document()
    for line in (
        "1. Період виконання: з 01.09.2026 по 31.05.2027",
        "Прізвище, власне ім’я, по батькові (за наявності): Шевченко Олена Петрівна",
        "Клас навчання: 6-А",
        "Рік навчання: 4-й",
    ):
        doc.add_paragraph(line)
    phone = doc.add_paragraph("Контактні номери телефонів батьків (інших законних представників) учениці: ")
    phone.add_run("+380(00-000-00-00)").underline = True
    doc.add_paragraph("5. Створення інклюзивного освітнього середовища, зокрема:")
    doc.add_paragraph("11. Погодження індивідуальної програми розвитку з батьками")
    doc.save(path)
    parsed = parse_document(path)
    assert parsed.fields["child_grade"] == "6-А"
    assert parsed.fields["year_of_study"] == "4"
    assert parsed.fields["parents_phones"] == "+380(00-000-00-00)"


def test_mapping_missing_required_and_invalid_date():
    spec = FieldSpec("birth", "Дата народження", "#birth", "date", True, {}, None, None, "input")
    with pytest.raises(DocumentParsingError):
        map_document({}, (spec,))
    with pytest.raises(DocumentParsingError):
        map_document({comparison_key(spec.docx_label): "31.02.2020"}, (spec,))
    result = map_document({comparison_key(spec.docx_label): "01.02.2020"}, (spec,))
    assert result[0].expected == "2020-02-01"


def test_state_machine_rejects_skipped_confirmation():
    machine = StateMachine()
    machine.advance(State.INPUT_VALIDATED)
    with pytest.raises(ValueError):
        machine.advance(State.FORM_POPULATED)
    machine.advance(State.ERROR)
    with pytest.raises(ValueError):
        machine.advance(State.DOCUMENT_FOUND)


def test_partial_real_site_selectors_cannot_start_automation():
    import json

    settings = json.loads(Path("config/settings.example.json").read_text(encoding="utf-8"))
    with pytest.raises(ConfigurationError, match="success_indicator"):
        require_selectors(settings["selectors"])


def test_appendix_three_filename_matches_exact_pupil(tmp_path: Path):
    path = tmp_path / "ІПР (Додаток 3)_Забашта Поліна.docx"
    path.touch()
    patterns = ("{Surname} {Name}.docx", "{Name} {Surname}.docx")
    assert find_document(tmp_path, parse_child_name("Поліна Забашта"), patterns) == path
    with pytest.raises(DocumentNotFound):
        find_document(tmp_path, parse_child_name("Інша Забашта"), patterns)


def test_school_ipr_split_patronymic_and_ordinal_year(tmp_path: Path):
    doc = Document()
    for line in (
        "1. Період виконання 01.09.2026-28.05.2027",
        "прізвище, власне ім’я, по батькові (за наявності) Забашта Поліна",
        "Євгенівна",
        "контактні номери телефонів батьків (інших законних представників) учня 0000000000",
        "клас навчання 7-Б",
        "рік навчання другий",
        "5. Створення інклюзивного освітнього середовища, зокрема:",
        "11. Погодження індивідуальної програми розвитку з батьками",
    ):
        doc.add_paragraph(line)
    path = tmp_path / "split.docx"
    doc.save(path)
    parsed = parse_document(path)
    assert parsed.fields["docx_child_full_name"] == "Забашта Поліна Євгенівна"
    assert parsed.fields["year_of_study"] == "2"
    assert parsed.fields["child_grade"] == "7-Б"
