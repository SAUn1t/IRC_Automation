"""Exact, local transfers from the 2026 school IPR DOCX to its matching portal form."""

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from .errors import DocumentParsingError
from .normalization import comparison_key, normalize_space


@dataclass(frozen=True)
class PortalField:
    name: str
    value: str
    kind: str
    row_label: str = ""


_ENVIRONMENT_ROWS = (
    ("Доступність", 1), ("Освітлення", 2), ("Рівень шуму", 3),
    ("Потреба в усамітненні, його періодичність і місце", 4),
    ("Інформаційні позначки", 5), ("Меблі", 6), ("Інше", 7),
    ("використання візуального розкладу", 8),
    ("використання індивідуальних позначок", 9),
    ("збільшення часу на виконання завдань", 10),
    ("збільшення обсягу допомоги", 11),
    ("чергування видів діяльності", 12),
    ("використання заохочень", 13),
    ("використання засобів концентрації уваги", 14),
    ("картки-підказки, картки-інструкції", 15),
    ("інше", 16),
    ("зміна змісту навчальних предметів", 17),
    ("зміна вимог щодо виконання завдань", 18),
    ("зміна підходів до проведення моніторингу", 19),
    ("інше", 20),
    ("Жестова мова", 21), ("Шрифт Брайля", 22),
    ("Календарна система", 23),
    ("Комунікаційна система обміну зображеннями", 24),
    ("Інше", 25),
)

_ASSESSMENT_ROWS = (
    "спрощення інструкцій", "гнучкість використання часу",
    "зменшення кількості завдань", "використання оптимальних форм оцінювання",
    "використання доступних форм викладу інформації",
    "використання допоміжних засобів", "уникнення графічних завдань",
    "уникнення зорових/слухових подразників",
)


def _subject_choice(value: str) -> str:
    choice = comparison_key(value.lstrip("+\uf07f☐□☑☒•●· ").strip())
    if choice == "адаптація":
        return "Адаптація"
    if choice == "модифікація":
        return "Модифікація"
    first_sentence = choice.split(".", 1)[0]
    if (first_sentence.startswith("адаптація ") and "не застосовується" not in first_sentence
            and re.search(r"\bмодифікаці\w*[^.]{0,200}\bне застосовується\.?$", choice)):
        return "Адаптація"
    return value


def _selection(text: str) -> tuple[bool, str] | None:
    value = normalize_space(text)
    label = value.lstrip("+\uf07f☐□☑☒ ")
    markers = value[:len(value) - len(label)]
    if not markers.strip():
        return None
    return "+" in markers or "☑" in markers or "☒" in markers, label.strip()


def _underlined_choice(paragraph: Paragraph) -> bool:
    """An underlined option label is selected; underlined boxes/lines alone are not."""
    return any(run.underline and any(char.isalpha() for char in run.text)
               for run in paragraph.runs)


def _assessment_choice(text: str) -> tuple[bool, str]:
    """Read mixed plus/bullet prefixes without mistaking a bullet for a label."""
    value = normalize_space(text)
    label = value.lstrip("+•●·\uf07f☐□ ")
    markers = value[:len(value) - len(label)]
    return "+" in markers, label.strip()


def _section(paragraphs: list[str], number: int) -> list[str]:
    starts = [i for i, text in enumerate(paragraphs) if re.match(rf"^{number}\.\s", text)]
    if len(starts) != 1:
        return []
    start = starts[0]
    end = next((i for i in range(start + 1, len(paragraphs))
                if re.match(rf"^{number + 1}\.\s", paragraphs[i])), len(paragraphs))
    return paragraphs[start:end]


def _unique_suffix(paragraphs: list[str], prefix: str) -> str | None:
    matches = [text[len(prefix):].strip(" :") for text in paragraphs
               if comparison_key(text).startswith(comparison_key(prefix))]
    if len(matches) > 1:
        raise DocumentParsingError(f"Duplicate school IPR label: {prefix}")
    return matches[0] if matches and matches[0] else None


def extract_portal_fields(path: Path, base: dict[str, str]) -> tuple[PortalField, ...]:
    """Return only DOCX values with the same meaning in the current 11-section form."""
    doc = Document(path)
    paragraphs = [normalize_space(p.text) for p in doc.paragraphs]
    fields: list[PortalField] = []
    seen: set[str] = set()

    def add(name: str, value: str | None, kind: str = "text", row_label: str = "") -> None:
        if value is None or not value.strip():
            return
        if name in seen:
            raise DocumentParsingError(f"Duplicate portal field mapping: {name}")
        seen.add(name)
        fields.append(PortalField(name, value.strip(), kind, row_label))

    for name, kind in (("execution_period", "text"), ("year_of_study", "select"),
                       ("child_grade", "text"), ("parents_phones", "text")):
        add(name, base[name], kind)

    for name, prefix, kind in (
        ("parent_address", "адреса задекларованого/зареєстрованого місця проживання (перебування)", "textarea"),
        ("child_enrollment_date", "дата зарахування до закладу загальної середньої освіти", "text"),
    ):
        add(name, _unique_suffix(paragraphs, prefix), kind)

    section3 = _section(paragraphs, 3)
    if section3 and ":" in section3[0]:
        add("3[1][1]", section3[0].split(":", 1)[1], "textarea")

    section4 = _section(paragraphs, 4)
    if section4:
        match = re.search(r"протокол засідання від\s+(\d{1,2})\s+([а-щьюяіїєґ]+)\s+(20\d{2})\s*р\.\s*№\s*(\d+)", section4[0], re.I)
        if match:
            day, month, year, number = match.groups()
            add("protocol_date", day, "number")
            add("protocol_month", month, "text")
            add("protocol_year", year[2:], "number")
            add("protocol_number", number, "number")
        level_labels = ("перший", "другий", "третій", "четвертий", "п’ятий")
        if len(section4) >= 6:
            section4_start = paragraphs.index(section4[0])
            levels = []
            for i, expected in enumerate(level_labels, 1):
                marked = _selection(section4[i])
                label = marked[1] if marked else section4[i]
                if not comparison_key(label).startswith(comparison_key(expected)):
                    break
                checked = bool((marked and marked[0]) or _underlined_choice(doc.paragraphs[section4_start + i]))
                levels.append((i, checked))
            if len(levels) == 5 and sum(checked for _, checked in levels) == 1:
                add("supportLevel", str(next(i for i, checked in levels if checked)), "select")

    environment = next((table for table in doc.tables if table.rows and
                        comparison_key(table.cell(0, 0).text).startswith(comparison_key("Назва адаптації/модифікації освітнього середовища"))), None)
    if environment:
        candidates = []
        for row in environment.rows[1:]:
            label = normalize_space(row.cells[0].text)
            if any(comparison_key(label).startswith(comparison_key(expected)) for expected, _ in _ENVIRONMENT_ROWS):
                candidates.append((label, normalize_space(row.cells[1].text)))
        if len(candidates) != len(_ENVIRONMENT_ROWS):
            raise DocumentParsingError("School IPR environment row layout changed")
        for (label, value), (expected, number) in zip(candidates, _ENVIRONMENT_ROWS):
            if not comparison_key(label).startswith(comparison_key(expected)):
                raise DocumentParsingError("School IPR environment row order changed")
            add(f"5[{number}][1]", value, "textarea", expected)

    section5 = _section(paragraphs, 5)
    if section5:
        section5_start = paragraphs.index(section5[0])
        section5_paragraphs = doc.paragraphs[section5_start:section5_start + len(section5)]

        def choice_at(index: int) -> tuple[bool, str]:
            text = section5[index]
            marked = _selection(text)
            label = marked[1] if marked else text
            return bool((marked and marked[0]) or _underlined_choice(section5_paragraphs[index])), label

        def selected_after(label: str, count: int = 2) -> str | None:
            positions = [i for i, text in enumerate(section5)
                         if comparison_key(text).startswith(comparison_key(label))]
            if len(positions) != 1:
                return None
            indices = range(positions[0] + 1, positions[0] + 1 + count)
            if indices.stop > len(section5):
                return None
            choices = [choice_at(index) for index in indices]
            if sum(checked for checked, _ in choices) != 1:
                return None
            return next(value for checked, value in choices if checked)

        plan_positions = [i for i, text in enumerate(section5)
                          if comparison_key(text).startswith("наявність індивідуального навчального плану")]
        if len(plan_positions) == 1:
            plan_indices = range(plan_positions[0] + 1, plan_positions[0] + 3)
            plan_choices = [choice_at(index) for index in plan_indices] if plan_indices.stop <= len(section5) else []
            if ([comparison_key(label) for _, label in plan_choices] == ["так", "ні"]
                    and sum(checked for checked, _ in plan_choices) <= 1):
                # The portal defaults to "так"; an unmarked DOCX pair means "ні".
                add("5[26][1]", "1" if plan_choices[0][0] else "0", "radio")
        schedule = selected_after("здійснення освітнього процесу")
        if schedule:
            add("5[27][1]", schedule.rstrip(" _"), "select")
        aids = selected_after("забезпечення допоміжними засобами для навчання")
        if aids and comparison_key(aids) in ("так", "ні"):
            add("5[28][1]", "1" if comparison_key(aids) == "так" else "0", "radio")
            if comparison_key(aids) == "так":
                add("5[28][2]", _unique_suffix(section5, "перелік допоміжних засобів для навчання, яких потребує учень"), "textarea")

    tables = (
        ("Найменування навчального предмета", "5", 29, ("textarea", "select", "textarea")),
        ("Освітня галузь/", "6", 1, ("textarea", "textarea", "textarea", "textarea", "textarea")),
        ("Найменування заняття", "7", 1, ("textarea", "textarea", "textarea", "textarea")),
        ("Прізвище, власне ім’я,", "10", 1, ("textarea", "textarea")),
    )
    for first_header, section, base_row, kinds in tables:
        matches = [table for table in doc.tables if table.rows and
                   comparison_key(table.cell(0, 0).text).startswith(comparison_key(first_header))]
        if len(matches) > 1:
            raise DocumentParsingError(f"Duplicate school IPR table in section {section}")
        if not matches:
            continue
        table = matches[0]
        for index, row in enumerate(table.rows[1:], base_row):
            values = [normalize_space(cell.text) for cell in row.cells]
            if len(values) < len(kinds):
                raise DocumentParsingError(f"School IPR section {section} has too few columns")
            if not values[0]:
                continue
            for column, (value, kind) in enumerate(zip(values, kinds), 1):
                if section == "5" and column == 2:
                    value = _subject_choice(value)
                add(f"{section}[{index}][{column}]", value, kind)

    section8 = _section(paragraphs, 8)
    choices8 = []
    if section8:
        section8_start = paragraphs.index(section8[0])
        for index, value in enumerate(section8[1:], 1):
            if value:
                marked, label = _assessment_choice(value)
                choices8.append((marked or _underlined_choice(doc.paragraphs[section8_start + index]), label))
    if choices8:
        if len(choices8) < len(_ASSESSMENT_ROWS):
            raise DocumentParsingError(
                f"School IPR section 8 has {len(choices8)} choices; expected {len(_ASSESSMENT_ROWS)}"
            )
        for index, label in enumerate(_ASSESSMENT_ROWS):
            if not comparison_key(choices8[index][1]).startswith(comparison_key(label)):
                raise DocumentParsingError(f"School IPR section 8 choice {index + 1} has an unexpected label")
        for i, (checked, _) in enumerate(choices8[:8], 1):
            add(f"8[{i}][1]", "true" if checked else "false", "checkbox")
        if len(choices8) >= 9 and comparison_key(choices8[8][1]).startswith("інше"):
            checked, other_text = choices8[8]
            if not checked:
                add("8[9][0]", "false", "checkbox")
            else:
                detail = other_text[len("інше"):].strip(" _.:;")
                if detail:
                    add("8[9][0]", "true", "checkbox")
                    add("8[9][1]", detail, "textarea")

    signer = (_unique_suffix(paragraphs, "Я,") or "").rstrip(" ,.;")
    parents = next((text for text in paragraphs
                    if comparison_key(text).startswith(
                        "прізвище, власне ім’я, по батькові (за наявності) батьків")), "")
    if signer and comparison_key(signer) in comparison_key(parents):
        add("footer_parents_names", signer, "textarea")
    return tuple(fields)
