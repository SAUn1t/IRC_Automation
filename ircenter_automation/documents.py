import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph
from zipfile import BadZipFile

from .errors import DocumentNotFound, DocumentParsingError, MultipleDocumentsFound
from .normalization import ChildName, comparison_key, labeled_school_full_name, normalize_space


def find_document(folder: Path, child: ChildName, patterns: tuple[str, ...]) -> Path:
    if not folder.is_dir():
        raise DocumentNotFound("Configured DOCX folder does not exist")
    names = {comparison_key(pattern.format(Name=first, Surname=surname))
             for pattern in patterns
             for first, surname in ((child.first, child.surname), (child.surname, child.first))}
    matches = sorted((p for p in folder.iterdir() if p.is_file() and p.suffix.casefold() == ".docx" and comparison_key(p.name) in names), key=lambda p: p.name)
    if not matches:
        raise DocumentNotFound("No exact DOCX filename match")
    if len(matches) > 1:
        raise MultipleDocumentsFound(f"{len(matches)} DOCX files match configured patterns; resolve this manually")
    return matches[0]


@dataclass(frozen=True)
class ParsedDocument:
    fields: dict[str, str]
    headings: tuple[str, ...]
    template: str = "generic"


def inspect_document(path: Path) -> str:
    """Return a local, value-free description of the document layout."""
    try:
        doc = Document(path)
    except (OSError, ValueError, KeyError, TypeError, BadZipFile, PackageNotFoundError) as exc:
        raise DocumentParsingError("Cannot read or parse the DOCX") from exc
    paragraphs = [normalize_space(p.text) for p in doc.paragraphs]
    if (any(p.startswith("5. Створення інклюзивного освітнього середовища") for p in paragraphs)
            and any(p.startswith("11. Погодження індивідуальної програми розвитку") for p in paragraphs)):
        return "11-section school IPR; matches the current 11-section portal form"
    return "generic label/value DOCX"


def _parse_school_ipr_2026(doc) -> ParsedDocument:
    paragraphs = [normalize_space(p.text) for p in doc.paragraphs]
    fields: dict[str, str] = {}
    for key, label in (("execution_period", r"1\.\s*період виконання"),
                       ("child_grade", "клас навчання"),
                       ("year_of_study", "рік навчання")):
        matches = [match.group(1).strip() for value in paragraphs
                   if (match := re.fullmatch(rf"{label}(?:\s*:\s*|\s+)(.+)", value, re.IGNORECASE))]
        if len(matches) != 1 or not matches[0]:
            raise DocumentParsingError(f"Required school IPR field is missing or ambiguous: {key}")
        fields[comparison_key(key)] = matches[0]
    phones = [match.group(1).strip() for value in paragraphs
              if (match := re.fullmatch(
                  r"контактні номери телефонів батьків \(інших законних представників\) "
                  r"(?:учня|учениці)(?:\s*:\s*|\s+)(.+)", value, re.IGNORECASE))]
    if len(phones) != 1 or not phones[0]:
        raise DocumentParsingError("Required school IPR field is missing or ambiguous: parents_phones")
    fields[comparison_key("parents_phones")] = phones[0]
    # Some school templates put the patronymic in the following Word paragraph.
    # Join only a single short continuation; never join arbitrary adjacent fields.
    child_names = []
    for index, value in enumerate(paragraphs):
        name = labeled_school_full_name(value)
        if name is None and index + 1 < len(paragraphs):
            continuation = paragraphs[index + 1]
            if len(continuation.split()) == 1 and continuation.isalpha():
                name = labeled_school_full_name(f"{value} {continuation}")
        if name is not None:
            child_names.append(name)
    if len(child_names) != 1:
        raise DocumentParsingError("School IPR child name is missing or ambiguous")
    fields[comparison_key("docx_child_full_name")] = child_names[0]
    year_text = comparison_key(fields[comparison_key("year_of_study")])
    ordinal_years = {
        "перший": 1, "другий": 2, "третій": 3, "четвертий": 4,
        "п’ятий": 5, "п'ятий": 5, "шостий": 6, "сьомий": 7,
        "восьмий": 8, "дев’ятий": 9, "дев'ятий": 9,
        "десятий": 10, "одинадцятий": 11, "дванадцятий": 12,
    }
    if year_text in ordinal_years:
        fields[comparison_key("year_of_study")] = str(ordinal_years[year_text])
        return ParsedDocument(fields, (), "school_ipr_2026")
    year = re.fullmatch(r"(\d{1,2})(?:\s*[-‑–]\s*(?:й|ий|ій))?",
                        year_text, re.IGNORECASE)
    if not year or not 1 <= int(year.group(1)) <= 12:
        raise DocumentParsingError("School IPR year of study is invalid")
    fields[comparison_key("year_of_study")] = str(int(year.group(1)))
    return ParsedDocument(fields, (), "school_ipr_2026")


_PAIR = re.compile(r"^\s*([^:\n]{1,120})\s*:\s*(.*)$", re.DOTALL)


def _clean(value: str) -> str:
    value = value.replace("☑", "yes").replace("☒", "yes").replace("☐", "no")
    return "\n".join(normalize_space(line) for line in value.splitlines()).strip()


def _add(fields: dict[str, str], label: str, value: str) -> None:
    key = comparison_key(label)
    if not key:
        return
    if key in fields:
        raise DocumentParsingError(f"Duplicate DOCX label: {label}")
    fields[key] = _clean(value)


def parse_document(path: Path) -> ParsedDocument:
    try:
        if inspect_document(path).startswith("11-section school IPR"):
            return _parse_school_ipr_2026(Document(path))
        doc = Document(path)
        fields: dict[str, str] = {}
        headings: list[str] = []
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                if block.style and block.style.name.lower().startswith("heading"):
                    headings.append(_clean(text))
                else:
                    pair = _PAIR.match(text)
                    if pair:
                        _add(fields, pair.group(1), pair.group(2))
            elif isinstance(block, Table):
                for row in block.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if len(cells) >= 2 and cells[0]:
                        if len(cells) % 2 or any(not cells[i] for i in range(0, len(cells), 2)):
                            raise DocumentParsingError("Table row does not contain clear label/value pairs")
                        for i in range(0, len(cells), 2):
                            _add(fields, cells[i], cells[i + 1])
                    elif any(cells):
                        raise DocumentParsingError("Table row has a value without a label")
        if not fields:
            raise DocumentParsingError("No label/value pairs found in DOCX")
        return ParsedDocument(fields, tuple(headings))
    except DocumentParsingError:
        raise
    except (OSError, ValueError, KeyError, TypeError, BadZipFile, PackageNotFoundError) as exc:
        raise DocumentParsingError("Cannot read or parse the DOCX") from exc
